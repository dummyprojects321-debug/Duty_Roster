"""
core.py - database, business rules and PDF / Excel export for the Duty Roster app.

This module has NO Streamlit dependency, so every rule in it can be unit-tested
(see test_core.py).  app.py is only the user interface on top of it.

Data model (SQLite, one file):
    users        login accounts  (ADMIN + one account per station)
    staff        people          (name, designation, belt no, contact, active)
    lists        per-station dropdown lists: kind = 'beat' | 'shift'
    slots        per-station time slots (start - end)
    roster       one row per person per day: status, beat, shift, slot, remarks
    roster_day   a row exists  => that station's roster for that day is FINAL
    signature    per-station saved signature (label, officer name, image)
    audit        who did what, when
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import secrets
import shutil
import sqlite3
import tempfile
from contextlib import contextmanager
from datetime import date, datetime
from xml.sax.saxutils import escape

import pandas as pd

# =============================================================================
# CONFIG  (safe to edit)
# =============================================================================
HERE = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.environ.get("DUTY_DB_PATH", os.path.join(HERE, "duty_roster_v3.db"))
SEED_FILE = os.path.join(HERE, "DUTY_ROSTER_Standard.xlsx")  # used once, on an empty database
ORG_NAME = "PERA  |  DERB Gujranwala"
APP_TITLE = "Duty Roster"
ADMIN_USER = "ADMIN"

# First-run passwords. Every account MUST change its password at first login.
DEFAULT_ADMIN_PW = "admin123"
DEFAULT_STATIONS = {
    "GUJRANWALA CITY": "city2026",
    "GUJRANWALA SADDAR": "saddar2026",
    "KAMOKE": "kamoke2026",
    "NOWSHEHRA VIRKAN": "nowshehra2026",
    "WAZIRABAD": "wazirabad2026",
}
DEFAULT_SHIFTS = ["Morning", "Evening", "Night"]
DEFAULT_SLOTS = [("09:00", "17:00"), ("13:00", "21:00"), ("21:00", "06:00")]
DEFAULT_DESIGNATIONS = [
    "SENIOR SERGEANT", "SERGEANT", "ENFORCEMENT OFFICER",
    "INVESTIGATION OFFICER", "DRIVER", "MUHARRAR",
]
STATUSES = ["Duty", "Leave", "Off Day", "Training"]
SEED_ROSTER = True        # also create one starting roster (dated today) from the sample file
PDF_VERSION = 3           # bump when the PDF layout changes (invalidates cached PDFs)

PHONE_RE = re.compile(r"^03\d{2}-\d{7}$")


# =============================================================================
# SMALL HELPERS
# =============================================================================
def _s(v) -> str:
    """Anything -> clean string ('' for None / NaN)."""
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v).strip()


def nice(station: str) -> str:
    return _s(station).title()


def norm_name(v) -> str:
    return " ".join(_s(v).split()).title()


def norm_belt(v) -> str:
    s = re.sub(r"\s+", "", _s(v)).upper()
    return s[:-2] if re.fullmatch(r"\d+\.0", s) else s


def norm_designation(v) -> str:
    return " ".join(_s(v).split()).upper()


def norm_phone(v) -> str:
    s = _s(v)
    d = re.sub(r"\D", "", s)
    if len(d) == 12 and d.startswith("923"):
        d = "0" + d[2:]
    if len(d) == 10 and d.startswith("3"):      # Excel dropped the leading zero
        d = "0" + d
    if len(d) == 11 and d.startswith("03"):
        return f"{d[:4]}-{d[4:]}"
    return s


def phone_ok(p: str) -> bool:
    return bool(PHONE_RE.match(_s(p)))


def to_hhmm(v) -> str:
    if _s(v) == "":
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    s = _s(v)
    for f in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M:%S %p"):
        try:
            return datetime.strptime(s, f).strftime("%H:%M")
        except ValueError:
            pass
    return ""


def fmt12(v) -> str:
    s = to_hhmm(v)
    return datetime.strptime(s, "%H:%M").strftime("%I:%M %p").lstrip("0") if s else ""


def slot_label(start, end) -> str:
    return f"{fmt12(start)} \u2013 {fmt12(end)}"


def slot_minutes(label: str) -> int:
    m = re.search(r"(\d{1,2}):(\d{2})\s*([AaPp][Mm])", _s(label))
    if not m:
        return 99999
    h, mi, ap = int(m.group(1)) % 12, int(m.group(2)), m.group(3).upper()
    return (h + (12 if ap == "PM" else 0)) * 60 + mi


def long_date(iso: str) -> str:
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%A, %d %B %Y")


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def today_iso() -> str:
    return date.today().isoformat()


# =============================================================================
# DATABASE
# =============================================================================
SCHEMA = """
CREATE TABLE IF NOT EXISTS users(
    username TEXT PRIMARY KEY, role TEXT NOT NULL, salt TEXT, pw_hash TEXT,
    must_change INTEGER DEFAULT 1, created_at TEXT);
CREATE TABLE IF NOT EXISTS staff(
    id INTEGER PRIMARY KEY AUTOINCREMENT, station TEXT NOT NULL, name TEXT NOT NULL,
    designation TEXT DEFAULT '', belt_no TEXT DEFAULT '', contact TEXT DEFAULT '',
    active INTEGER DEFAULT 1, created_at TEXT);
CREATE INDEX IF NOT EXISTS ix_staff_station ON staff(station, active);
CREATE UNIQUE INDEX IF NOT EXISTS ux_staff_belt ON staff(belt_no) WHERE belt_no <> '';
CREATE TABLE IF NOT EXISTS lists(
    id INTEGER PRIMARY KEY AUTOINCREMENT, station TEXT NOT NULL, kind TEXT NOT NULL,
    name TEXT NOT NULL COLLATE NOCASE, UNIQUE(station, kind, name));
CREATE TABLE IF NOT EXISTS slots(
    id INTEGER PRIMARY KEY AUTOINCREMENT, station TEXT NOT NULL, t_start TEXT, t_end TEXT,
    label TEXT NOT NULL, UNIQUE(station, label));
CREATE TABLE IF NOT EXISTS roster(
    id INTEGER PRIMARY KEY AUTOINCREMENT, station TEXT NOT NULL, roster_date TEXT NOT NULL,
    staff_id INTEGER NOT NULL REFERENCES staff(id), status TEXT NOT NULL DEFAULT 'Duty',
    beat TEXT DEFAULT '', shift TEXT DEFAULT '', slot TEXT DEFAULT '', remarks TEXT DEFAULT '',
    updated_by TEXT, updated_at TEXT, UNIQUE(roster_date, staff_id));
CREATE INDEX IF NOT EXISTS ix_roster_day ON roster(roster_date, station);
CREATE TABLE IF NOT EXISTS roster_day(
    station TEXT, roster_date TEXT, by_user TEXT, at TEXT, PRIMARY KEY(station, roster_date));
CREATE TABLE IF NOT EXISTS signature(
    station TEXT PRIMARY KEY, label TEXT, officer TEXT, image BLOB);
CREATE TABLE IF NOT EXISTS audit(
    id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT, user TEXT, action TEXT, detail TEXT);
"""


@contextmanager
def db():
    conn = sqlite3.connect(DB_FILE, timeout=20)
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def audit(user: str, action: str, detail: str = "", conn=None):
    sql = "INSERT INTO audit(ts, user, action, detail) VALUES (?,?,?,?)"
    args = (now_iso(), user, action, detail)
    if conn is not None:
        conn.execute(sql, args)
    else:
        with db() as c:
            c.execute(sql, args)


# ---- passwords ---------------------------------------------------------------
def _hash(pw: str, salt_hex: str) -> str:
    return hashlib.pbkdf2_hmac("sha256", pw.encode(), bytes.fromhex(salt_hex), 150_000).hex()


def _set_pw(c, username: str, role: str, pw: str, must_change: bool):
    salt = secrets.token_hex(16)
    c.execute(
        "INSERT INTO users(username, role, salt, pw_hash, must_change, created_at) VALUES (?,?,?,?,?,?) "
        "ON CONFLICT(username) DO UPDATE SET salt=excluded.salt, pw_hash=excluded.pw_hash, "
        "must_change=excluded.must_change",
        (username, role, salt, _hash(pw, salt), int(must_change), now_iso()),
    )


def authenticate(username: str, pw: str):
    with db() as c:
        row = c.execute(
            "SELECT username, role, salt, pw_hash, must_change FROM users WHERE username=?",
            (_s(username).upper(),)).fetchone()
    if not row or not secrets.compare_digest(_hash(pw or "", row[2]), row[3]):
        return None
    user = {"username": row[0], "role": row[1], "must_change": bool(row[4]),
            "station": row[0] if row[1] == "station" else None}
    audit(user["username"], "login")
    return user


def change_password(username: str, new_pw: str, by: str | None = None):
    with db() as c:
        role = c.execute("SELECT role FROM users WHERE username=?", (username,)).fetchone()
        if not role:
            raise ValueError("Unknown user")
        # a password set by someone else (admin reset) must be changed at next login
        _set_pw(c, username, role[0], new_pw, must_change=bool(by and by != username))
        audit(by or username, "password_change", username, c)


# ---- stations ----------------------------------------------------------------
def list_stations() -> list[str]:
    with db() as c:
        return [r[0] for r in c.execute("SELECT username FROM users WHERE role='station' ORDER BY rowid")]


def _ensure_defaults(c, station: str):
    if not c.execute("SELECT 1 FROM lists WHERE station=? AND kind='shift'", (station,)).fetchone():
        for n in DEFAULT_SHIFTS:
            c.execute("INSERT OR IGNORE INTO lists(station, kind, name) VALUES (?,?,?)", (station, "shift", n))
    if not c.execute("SELECT 1 FROM slots WHERE station=?", (station,)).fetchone():
        for s, e in DEFAULT_SLOTS:
            c.execute("INSERT OR IGNORE INTO slots(station, t_start, t_end, label) VALUES (?,?,?,?)",
                      (station, s, e, slot_label(s, e)))


def add_station(name: str, password: str, by: str):
    name = " ".join(_s(name).split()).upper()
    if not name:
        return False, "Enter a station name."
    if len(password) < 6:
        return False, "Password must be at least 6 characters."
    with db() as c:
        if c.execute("SELECT 1 FROM users WHERE username=?", (name,)).fetchone():
            return False, f"'{name}' already exists."
        _set_pw(c, name, "station", password, True)
        _ensure_defaults(c, name)
        audit(by, "add_station", name, c)
    return True, f"Station {nice(name)} created. It must change the password at first login."


# =============================================================================
# STAFF
# =============================================================================
def designations() -> list[str]:
    with db() as c:
        used = [r[0] for r in c.execute(
            "SELECT DISTINCT designation FROM staff WHERE designation<>'' ORDER BY designation")]
    return DEFAULT_DESIGNATIONS + [d for d in used if d not in DEFAULT_DESIGNATIONS]


def staff_df(station: str, active_only: bool = True) -> pd.DataFrame:
    q = "SELECT id, name, designation, belt_no, contact, active FROM staff WHERE station=?"
    if active_only:
        q += " AND active=1"
    q += " ORDER BY name COLLATE NOCASE"
    with db() as c:
        df = pd.read_sql_query(q, c, params=(station,))
    df["active"] = df["active"].astype(bool)
    return df


def _belt_owner(c, belt: str, exclude_id: int | None = None):
    if not belt:
        return None
    return c.execute("SELECT station, name FROM staff WHERE belt_no=? AND id<>?",
                     (belt, exclude_id or -1)).fetchone()


def _insert_staff(c, station, name, designation, belt, contact):
    """Insert one person; returns new id, or None if the belt number is already used."""
    if _belt_owner(c, belt):
        return None
    cur = c.execute(
        "INSERT INTO staff(station, name, designation, belt_no, contact, active, created_at) "
        "VALUES (?,?,?,?,?,1,?)", (station, name, designation, belt, contact, now_iso()))
    return cur.lastrowid


def add_staff(station, name, designation, belt, contact, by):
    name, des, belt, ph = norm_name(name), norm_designation(designation), norm_belt(belt), norm_phone(contact)
    if not name:
        return False, "Name is required."
    with db() as c:
        owner = _belt_owner(c, belt)
        if owner:
            return False, f"Belt no. {belt} already belongs to {owner[1]} ({nice(owner[0])})."
        if not belt and c.execute(
                "SELECT 1 FROM staff WHERE station=? AND belt_no='' AND lower(name)=lower(?)",
                (station, name)).fetchone():
            return False, f"{name} is already in your list."
        _insert_staff(c, station, name, des, belt, ph)
        audit(by, "add_staff", f"{name} {belt}", c)
    msg = f"Added {name}."
    if ph and not phone_ok(ph):
        msg += " (Check the contact number format: 03XX-XXXXXXX.)"
    return True, msg


def update_staff(station, df: pd.DataFrame, by):
    """Apply edits made in the staff table. Returns (n_updated, [error messages])."""
    n, errors = 0, []
    with db() as c:
        for r in df.to_dict("records"):
            sid = int(r["id"])
            if not c.execute("SELECT 1 FROM staff WHERE id=? AND station=?", (sid, station)).fetchone():
                continue
            name, des = norm_name(r["name"]), norm_designation(r["designation"])
            belt, ph = norm_belt(r["belt_no"]), norm_phone(r["contact"])
            if not name:
                errors.append("A row has no name - skipped.")
                continue
            owner = _belt_owner(c, belt, exclude_id=sid)
            if owner:
                errors.append(f"{name}: belt {belt} already belongs to {owner[1]}.")
                continue
            cur = c.execute(
                "UPDATE staff SET name=?, designation=?, belt_no=?, contact=?, active=? "
                "WHERE id=? AND (name<>? OR designation<>? OR belt_no<>? OR contact<>? OR active<>?)",
                (name, des, belt, ph, int(bool(r["active"])), sid,
                 name, des, belt, ph, int(bool(r["active"]))))
            n += cur.rowcount
        if n:
            audit(by, "update_staff", f"{station}: {n} rows", c)
    return n, errors


def transfer_staff(staff_id: int, new_station: str, by: str):
    """Move a person to another station. Their upcoming duties at the old station are removed."""
    with db() as c:
        row = c.execute("SELECT name, station FROM staff WHERE id=?", (staff_id,)).fetchone()
        if not row:
            return False, "Person not found."
        if row[1] == new_station:
            return False, "Already at that station."
        c.execute("UPDATE staff SET station=? WHERE id=?", (new_station, staff_id))
        gone = c.execute("DELETE FROM roster WHERE staff_id=? AND roster_date>=?",
                         (staff_id, today_iso())).rowcount
        audit(by, "transfer", f"{row[0]}: {row[1]} -> {new_station}", c)
    return True, f"{row[0]} moved to {nice(new_station)}. {gone} upcoming duty entr{'y' if gone == 1 else 'ies'} removed."


# ---- bulk import --------------------------------------------------------------
_ALIASES = {
    "name": {"name", "staffname", "employeename", "fullname", "employee"},
    "designation": {"designation", "rank", "post", "title"},
    "belt_no": {"belt", "beltno", "beltnumber", "badge", "badgeno", "badgenumber", "serviceno"},
    "contact": {"contact", "contactno", "contactnumber", "phone", "phoneno", "phonenumber",
                "mobile", "mobileno", "cell", "cellno"},
}
IMPORT_COLS = ["name", "designation", "belt_no", "contact"]


def standardize_import(df: pd.DataFrame) -> pd.DataFrame:
    """Map whatever headers the file has onto name / designation / belt_no / contact."""
    mapping = {}
    for col in df.columns:
        key = re.sub(r"[^a-z0-9]", "", str(col).lower())
        for target, names in _ALIASES.items():
            if key in names and target not in mapping.values():
                mapping[col] = target
    if "name" in mapping.values():
        out = df.rename(columns=mapping)
    else:  # unknown headers -> trust the column order: name, designation, belt, contact
        out = df.iloc[:, :4].copy()
        out.columns = IMPORT_COLS[: out.shape[1]]
    for col in IMPORT_COLS:
        if col not in out.columns:
            out[col] = ""
    return out[IMPORT_COLS].astype(object).where(out[IMPORT_COLS].notna(), "")


def parse_pasted(text: str) -> pd.DataFrame:
    rows = []
    for line in text.splitlines():
        if not line.strip():
            continue
        sep = "\t" if "\t" in line else ("|" if "|" in line else ",")
        cells = [x.strip() for x in line.split(sep)][:4]
        rows.append(cells + [""] * (4 - len(cells)))
    df = pd.DataFrame(rows, columns=IMPORT_COLS)
    if len(df) and re.sub(r"[^a-z]", "", df.iloc[0]["name"].lower()) in _ALIASES["name"]:
        df = df.iloc[1:]
    return df.reset_index(drop=True)


def read_import_file(file) -> pd.DataFrame:
    name = getattr(file, "name", "").lower()
    raw = pd.read_csv(file, dtype=str) if name.endswith(".csv") else pd.read_excel(file, dtype=str)
    return standardize_import(raw)


def prepare_import(df: pd.DataFrame, station: str) -> pd.DataFrame:
    """Clean + validate rows. Adds a 'status' column; only 'New' / 'Check phone' rows get imported."""
    seen_belts, seen_names, out = set(), set(), []
    with db() as c:
        belts = {r[0]: r[1] for r in c.execute("SELECT belt_no, station FROM staff WHERE belt_no<>''")}
        names = {r[0] for r in c.execute(
            "SELECT lower(name) FROM staff WHERE station=? AND belt_no=''", (station,))}
    for r in df.to_dict("records"):
        name, des = norm_name(r.get("name")), norm_designation(r.get("designation"))
        belt, ph = norm_belt(r.get("belt_no")), norm_phone(r.get("contact"))
        if not name:
            status = "Name missing"
        elif belt and belt in belts:
            status = "Already exists"
        elif belt and belt in seen_belts:
            status = "Duplicate in file"
        elif not belt and (name.lower() in names or name.lower() in seen_names):
            status = "Duplicate in file" if name.lower() in seen_names else "Already exists"
        elif ph and not phone_ok(ph):
            status = "Check phone"
        else:
            status = "New"
        if status in ("New", "Check phone"):
            (seen_belts.add(belt) if belt else seen_names.add(name.lower()))
        out.append({"name": name, "designation": des, "belt_no": belt, "contact": ph, "status": status})
    return pd.DataFrame(out, columns=IMPORT_COLS + ["status"])


def import_staff(station: str, prepared: pd.DataFrame, by: str) -> int:
    ok = prepared[prepared["status"].isin(["New", "Check phone"])]
    n = 0
    with db() as c:
        for r in ok.to_dict("records"):
            if _insert_staff(c, station, r["name"], r["designation"], r["belt_no"], r["contact"]):
                n += 1
        audit(by, "import_staff", f"{station}: {n} people", c)
    return n


def import_template_bytes() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(["Name", "Designation", "Belt No", "Contact No"])
    ws.append(["Muhammad Ali", "SERGEANT", "SGT-1001", "0300-1234567"])
    ws.append(["Ahmed Raza", "ENFORCEMENT OFFICER", "EO-215", "0321-7654321"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="12263F")
        cell.alignment = Alignment(horizontal="center")
    for col, w in zip("ABCD", (28, 26, 14, 16)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# DROPDOWN LISTS: beats, shifts, time slots
# =============================================================================
def get_list(kind: str, station: str) -> list[str]:
    with db() as c:
        return [r[0] for r in c.execute(
            "SELECT name FROM lists WHERE station=? AND kind=? ORDER BY id", (station, kind))]


def _add_list(c, kind, station, names) -> int:
    n = 0
    for nm in names:
        nm = " ".join(_s(nm).split())
        if nm:
            n += c.execute("INSERT OR IGNORE INTO lists(station, kind, name) VALUES (?,?,?)",
                           (station, kind, nm)).rowcount
    return n


def add_to_list(kind, station, names, by) -> int:
    with db() as c:
        n = _add_list(c, kind, station, names)
        if n:
            audit(by, f"add_{kind}", f"{station}: {n}", c)
    return n


def remove_from_list(kind, station, names, by) -> int:
    with db() as c:
        n = sum(c.execute("DELETE FROM lists WHERE station=? AND kind=? AND name=?",
                          (station, kind, nm)).rowcount for nm in names)
        if n:
            audit(by, f"remove_{kind}", f"{station}: {n}", c)
    return n


def get_slots(station: str) -> list[str]:
    with db() as c:
        rows = [r[0] for r in c.execute("SELECT label FROM slots WHERE station=?", (station,))]
    return sorted(rows, key=lambda x: (slot_minutes(x), x))


def add_slot(station, start, end, by):
    s, e = to_hhmm(start), to_hhmm(end)
    if not s or not e:
        return False, "Pick a start and an end time."
    if s == e:
        return False, "Start and end time cannot be the same."
    with db() as c:
        n = c.execute("INSERT OR IGNORE INTO slots(station, t_start, t_end, label) VALUES (?,?,?,?)",
                      (station, s, e, slot_label(s, e))).rowcount
        audit(by, "add_slot", slot_label(s, e), c)
    return (True, f"Added {slot_label(s, e)}.") if n else (False, "That time slot already exists.")


def remove_slots(station, labels, by) -> int:
    with db() as c:
        n = sum(c.execute("DELETE FROM slots WHERE station=? AND label=?", (station, lb)).rowcount
                for lb in labels)
        if n:
            audit(by, "remove_slot", f"{station}: {n}", c)
    return n


# =============================================================================
# ROSTER
# =============================================================================
def _clean_fields(status, beat, shift, slot):
    status = status if status in STATUSES else "Duty"
    beat, shift, slot = " ".join(_s(beat).split()), " ".join(_s(shift).split()), _s(slot)
    if status != "Duty":
        beat = shift = slot = ""
    return status, beat, shift, slot


def _upsert(c, station, iso, sid, status, beat, shift, slot, remarks, by):
    """remarks=None keeps whatever remark the person already has."""
    cols = "station, roster_date, staff_id, status, beat, shift, slot, remarks, updated_by, updated_at"
    upd = ("station=excluded.station, status=excluded.status, beat=excluded.beat, shift=excluded.shift, "
           "slot=excluded.slot, updated_by=excluded.updated_by, updated_at=excluded.updated_at")
    if remarks is not None:
        upd += ", remarks=excluded.remarks"
    c.execute(
        f"INSERT INTO roster({cols}) VALUES (?,?,?,?,?,?,?,?,?,?) "
        f"ON CONFLICT(roster_date, staff_id) DO UPDATE SET {upd}",
        (station, iso, sid, status, beat, shift, slot, remarks or "", by, now_iso()))


def _reopen(c, station, iso):
    c.execute("DELETE FROM roster_day WHERE station=? AND roster_date=?", (station, iso))


def roster_grid(station: str, iso: str) -> pd.DataFrame:
    """Every active person of the station + whatever is already assigned on that date."""
    with db() as c:
        df = pd.read_sql_query(
            """SELECT s.id AS staff_id, s.name, s.designation, s.belt_no, s.contact,
                      COALESCE(r.status,'Duty') AS status, COALESCE(r.beat,'') AS beat,
                      COALESCE(r.shift,'') AS shift, COALESCE(r.slot,'') AS slot,
                      COALESCE(r.remarks,'') AS remarks
               FROM staff s
               LEFT JOIN roster r ON r.staff_id=s.id AND r.roster_date=? AND r.station=?
               WHERE (s.station=? AND s.active=1) OR r.id IS NOT NULL
               ORDER BY s.name COLLATE NOCASE""",
            c, params=(iso, station, station))
    return df.fillna("")


def save_roster(station: str, iso: str, df: pd.DataFrame, by: str) -> dict:
    """Persist the grid. Rows with status Duty and nothing chosen are 'unassigned' (removed)."""
    saved = cleared = 0
    warnings, changed = [], False
    with db() as c:
        existing = {r[0]: tuple(r[1:]) for r in c.execute(
            "SELECT staff_id, status, beat, shift, slot, remarks FROM roster WHERE roster_date=? AND station=?",
            (iso, station))}
        mine = {r[0] for r in c.execute("SELECT id FROM staff WHERE station=?", (station,))} | set(existing)
        for r in df.to_dict("records"):
            sid = int(r["staff_id"])
            if sid not in mine:
                continue
            status, beat, shift, slot = _clean_fields(r.get("status"), r.get("beat"), r.get("shift"), r.get("slot"))
            rem = " ".join(_s(r.get("remarks")).split())
            if status == "Duty" and not (beat or shift or slot):
                if sid in existing:
                    c.execute("DELETE FROM roster WHERE roster_date=? AND staff_id=?", (iso, sid))
                    cleared += 1
                    changed = True
                continue
            if status == "Duty":
                missing = [n for n, v in (("beat", beat), ("shift", shift), ("time slot", slot)) if not v]
                if missing:
                    warnings.append(f"{_s(r.get('name'))}: {', '.join(missing)} not set")
            if existing.get(sid) != (status, beat, shift, slot, rem):
                _upsert(c, station, iso, sid, status, beat, shift, slot, rem, by)
                changed = True
            saved += 1
        if changed:
            _reopen(c, station, iso)
            audit(by, "save_roster", f"{station} {iso}: {saved} entries", c)
    return {"saved": saved, "cleared": cleared, "warnings": warnings, "changed": changed}


def assign(station, iso, staff_ids, beat, shift, slot, by) -> int:
    """Give several people the same beat / shift / slot in one go (new beat / shift names are remembered)."""
    status, beat, shift, slot = _clean_fields("Duty", beat, shift, slot)
    n = 0
    with db() as c:
        if beat:
            _add_list(c, "beat", station, [beat])
        if shift:
            _add_list(c, "shift", station, [shift])
        mine = {r[0] for r in c.execute("SELECT id FROM staff WHERE station=? AND active=1", (station,))}
        for sid in staff_ids:
            if int(sid) in mine:
                _upsert(c, station, iso, int(sid), status, beat, shift, slot, None, by)
                n += 1
        if n:
            _reopen(c, station, iso)
            audit(by, "assign", f"{station} {iso}: {n} people -> {beat}", c)
    return n


def roster_dates(station: str, before: str | None = None) -> list[str]:
    q, args = "SELECT DISTINCT roster_date FROM roster WHERE station=?", [station]
    if before:
        q += " AND roster_date<?"
        args.append(before)
    q += " ORDER BY roster_date DESC LIMIT 60"
    with db() as c:
        return [r[0] for r in c.execute(q, args)]


def copy_roster(station, src_iso, dst_iso, by, replace=True) -> int:
    with db() as c:
        rows = c.execute(
            """SELECT r.staff_id, r.status, r.beat, r.shift, r.slot, r.remarks FROM roster r
               JOIN staff s ON s.id=r.staff_id
               WHERE r.station=? AND r.roster_date=? AND s.active=1 AND s.station=?""",
            (station, src_iso, station)).fetchall()
        if replace:
            c.execute("DELETE FROM roster WHERE station=? AND roster_date=?", (station, dst_iso))
        have = {r[0] for r in c.execute(
            "SELECT staff_id FROM roster WHERE station=? AND roster_date=?", (station, dst_iso))}
        n = 0
        for sid, status, beat, shift, slot, rem in rows:
            if sid in have:
                continue
            _upsert(c, station, dst_iso, sid, status, beat, shift, slot, rem, by)
            n += 1
        _reopen(c, station, dst_iso)
        audit(by, "copy_roster", f"{station} {src_iso} -> {dst_iso}: {n}", c)
    return n


def clear_roster(station, iso, by) -> int:
    with db() as c:
        n = c.execute("DELETE FROM roster WHERE station=? AND roster_date=?", (station, iso)).rowcount
        _reopen(c, station, iso)
        audit(by, "clear_roster", f"{station} {iso}: {n}", c)
    return n


def set_final(station, iso, final: bool, by):
    with db() as c:
        if final:
            c.execute("INSERT OR REPLACE INTO roster_day(station, roster_date, by_user, at) VALUES (?,?,?,?)",
                      (station, iso, by, now_iso()))
        else:
            _reopen(c, station, iso)
        audit(by, "finalize" if final else "reopen", f"{station} {iso}", c)


def is_final(station, iso) -> bool:
    with db() as c:
        return bool(c.execute("SELECT 1 FROM roster_day WHERE station=? AND roster_date=?",
                              (station, iso)).fetchone())


def _station_order(stations=None):
    allst = list_stations()
    return allst if stations is None else [s for s in allst if s in set(stations)]


def roster_rows(iso: str, stations=None) -> pd.DataFrame:
    """All roster entries of a day, ready to print: sorted by station, shift, time slot, beat, name."""
    stations = _station_order(stations)
    cols = ["station", "staff_id", "name", "designation", "belt_no", "contact", "active",
            "status", "beat", "shift", "slot", "remarks"]
    if not stations:
        return pd.DataFrame(columns=cols)
    ph = ",".join("?" * len(stations))
    with db() as c:
        df = pd.read_sql_query(
            f"""SELECT r.station, r.staff_id, s.name, s.designation, s.belt_no, s.contact, s.active,
                       r.status, r.beat, r.shift, r.slot, r.remarks
                FROM roster r JOIN staff s ON s.id=r.staff_id
                WHERE r.roster_date=? AND r.station IN ({ph})""", c, params=[iso, *stations])
        shift_rank = {(st, nm.lower()): i for st in stations for i, nm in enumerate(get_list("shift", st))}
    df = df.fillna("")
    if df.empty:
        return df
    df["_st"] = df["station"].map({s: i for i, s in enumerate(stations)})
    df["_duty"] = (df["status"] != "Duty").astype(int)
    df["_sh"] = [shift_rank.get((a, b.lower()), 999) for a, b in zip(df["station"], df["shift"])]
    df["_t"] = df["slot"].map(slot_minutes)
    df["_b"] = df["beat"].str.lower()
    df["_n"] = df["name"].str.lower()
    df = df.sort_values(["_st", "_duty", "_sh", "_t", "_b", "_n"])
    return df.drop(columns=[c for c in df.columns if c.startswith("_")]).reset_index(drop=True)


def roster_state(iso: str) -> dict:
    """station -> 'final' | 'draft' for stations that have any entry on that date."""
    with db() as c:
        have = {r[0] for r in c.execute("SELECT DISTINCT station FROM roster WHERE roster_date=?", (iso,))}
        fin = {r[0] for r in c.execute("SELECT station FROM roster_day WHERE roster_date=?", (iso,))}
    return {s: ("final" if s in fin else "draft") for s in have | fin}


def day_summary(iso: str, stations=None) -> pd.DataFrame:
    stations = _station_order(stations)
    rows = roster_rows(iso, stations)
    state = roster_state(iso)
    with db() as c:
        counts = dict(c.execute("SELECT station, COUNT(*) FROM staff WHERE active=1 GROUP BY station"))
    out = []
    for stn in stations:
        r = rows[rows["station"] == stn]
        duty = r[r["status"] == "Duty"]
        total = int(counts.get(stn, 0))
        placed = int((r["active"] == 1).sum())
        out.append({
            "station": stn, "staff": total, "on_duty": len(duty),
            "leave_off": int((r["status"] != "Duty").sum()),
            "unassigned": max(total - placed, 0),
            "beats": int(duty["beat"].replace("", pd.NA).dropna().nunique()),
            "state": state.get(stn, "none"),
        })
    return pd.DataFrame(out, columns=["station", "staff", "on_duty", "leave_off", "unassigned", "beats", "state"])


def track(query: str, d1: str, d2: str, stations=None) -> pd.DataFrame:
    """History of where people were deputed. query matches name / belt / contact."""
    stations = _station_order(stations)
    if not stations:
        return pd.DataFrame()
    ph, q = ",".join("?" * len(stations)), f"%{_s(query)}%"
    with db() as c:
        return pd.read_sql_query(
            f"""SELECT r.roster_date AS date, r.station, s.name, s.belt_no, s.designation,
                       r.status, r.beat, r.shift, r.slot, r.remarks
                FROM roster r JOIN staff s ON s.id=r.staff_id
                WHERE r.roster_date BETWEEN ? AND ? AND r.station IN ({ph})
                  AND (s.name LIKE ? OR s.belt_no LIKE ? OR s.contact LIKE ?)
                ORDER BY r.roster_date DESC, s.name COLLATE NOCASE""",
            c, params=[d1, d2, *stations, q, q, q]).fillna("")


def beat_coverage(station: str, iso: str) -> pd.DataFrame:
    """Every beat of the station with the number of people on it (0 = uncovered)."""
    rows = roster_rows(iso, [station])
    duty = rows[rows["status"] == "Duty"]
    counts = duty.groupby("beat").size().to_dict()
    beats = get_list("beat", station)
    extra = [b for b in counts if b and b not in beats]
    return pd.DataFrame({"beat": beats + extra, "people": [counts.get(b, 0) for b in beats + extra]})


# =============================================================================
# SIGNATURE
# =============================================================================
def default_label(station: str) -> str:
    return f"SDEO ({station})"


def get_signature(station: str) -> dict:
    with db() as c:
        row = c.execute("SELECT label, officer, image FROM signature WHERE station=?", (station,)).fetchone()
    label, officer, img = (row or (None, None, None))
    return {"label": label or default_label(station), "officer": officer or "", "image": img}


def save_signature(station, label, officer, by):
    label = " ".join(_s(label).split()) or default_label(station)
    with db() as c:
        c.execute("INSERT INTO signature(station, label, officer) VALUES (?,?,?) "
                  "ON CONFLICT(station) DO UPDATE SET label=excluded.label, officer=excluded.officer",
                  (station, label, " ".join(_s(officer).split())))
        audit(by, "save_signature", station, c)


def set_signature_image(station, png: bytes | None, by):
    with db() as c:
        c.execute("INSERT INTO signature(station, label, officer, image) VALUES (?,?,?,?) "
                  "ON CONFLICT(station) DO UPDATE SET image=excluded.image",
                  (station, default_label(station), "", png))
        audit(by, "signature_image", "removed" if png is None else "saved", c)


def process_signature(raw: bytes) -> bytes:
    """Photo / scan of a signature -> small transparent PNG, cropped tight, paper background removed."""
    import numpy as np
    from PIL import Image, ImageOps
    im = ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert("RGBA")
    if im.width > 900:
        im = im.resize((900, max(1, round(im.height * 900 / im.width))), Image.LANCZOS)
    arr = np.array(im)
    if arr[..., 3].min() >= 250:  # opaque (photo / scan): make the white paper transparent
        lum = arr[..., :3] @ np.array([0.299, 0.587, 0.114])
        arr[..., 3] = np.clip((225 - lum) / 70 * 255, 0, 255).astype("uint8")
    im = Image.fromarray(arr)
    box = im.getchannel("A").point(lambda a: 255 if a > 25 else 0).getbbox()
    if not box:
        raise ValueError("No signature found in that image. Use dark ink on white paper.")
    pad = 6
    box = (max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, im.width), min(box[3] + pad, im.height))
    buf = io.BytesIO()
    im.crop(box).save(buf, "PNG", optimize=True)
    return buf.getvalue()


# =============================================================================
# EXPORT: PDF
# =============================================================================
def fingerprint(iso: str, stations) -> str:
    """Changes whenever anything that appears in the printout changes (used to cache PDFs)."""
    stations = list(stations)
    h = hashlib.sha1(f"v{PDF_VERSION}|{iso}|{stations}".encode())
    h.update(roster_rows(iso, stations).to_csv().encode())
    h.update(str(roster_state(iso)).encode())
    for s in stations:
        sig = get_signature(s)
        h.update(f"{sig['label']}|{sig['officer']}|{len(sig['image'] or b'')}".encode())
    return h.hexdigest()


def build_pdf(iso: str, stations, cover: bool = False) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Flowable, Frame, Image,
                                    PageBreak, PageTemplate, Paragraph, Spacer, Table)

    NAVY, ACCENT = colors.HexColor("#12263F"), colors.HexColor("#1AA3A3")
    INK, MUTED, LINE = colors.HexColor("#0F172A"), colors.HexColor("#64748B"), colors.HexColor("#D5DDE8")
    ZEBRA = colors.HexColor("#F1F5FA")
    BANDS = [  # keyword in shift name -> (background, text)
        ("morning", "#FFF3D1", "#8A5200"), ("mid", "#DDF3E4", "#17653A"), ("evening", "#E4E7FF", "#3338A0"),
        ("night", "#E2E8F0", "#0F172A"), ("sunday", "#FDE4F0", "#9D174D"),
    ]

    def band_colors(name):
        for key, bg, fg in BANDS:
            if key in name.lower():
                return colors.HexColor(bg), colors.HexColor(fg)
        return colors.HexColor("#E3EAF3"), colors.HexColor("#1E293B")

    stations = _station_order(stations)
    rows = roster_rows(iso, stations)
    summ = day_summary(iso, stations)
    generated = datetime.now().strftime("%d %b %Y, %I:%M %p").lstrip("0")
    title = "CONSOLIDATED DUTY ROSTER" if cover else "DAILY DUTY ROSTER"
    date_txt = long_date(iso)

    class Marker(Flowable):
        """Zero-size flowable that tells the page header which station is being printed."""
        def __init__(self, name):
            super().__init__()
            self.name = name

        def wrap(self, aw, ah):
            return 0, 0

        def draw(self):
            self.canv._cur_station = self.name

    class Numbered(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._saved, self._cur_station = [], ""

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total):
            w, h = self._pagesize
            self.setFillColor(NAVY)
            self.rect(0, h - 17 * mm, w, 17 * mm, stroke=0, fill=1)
            self.setFillColor(ACCENT)
            self.rect(0, h - 18.2 * mm, w, 1.2 * mm, stroke=0, fill=1)
            self.setFillColor(colors.white)
            self.setFont("Helvetica-Bold", 14)
            self.drawString(12 * mm, h - 9.5 * mm, title)
            self.setFont("Helvetica", 7.5)
            self.setFillColor(colors.HexColor("#B6C4D8"))
            self.drawString(12 * mm, h - 14 * mm, ORG_NAME)
            self.setFillColor(colors.white)
            self.setFont("Helvetica-Bold", 12)
            self.drawCentredString(w / 2 + 12 * mm, h - 10.5 * mm, self._cur_station)
            self.setFont("Helvetica-Bold", 9.5)
            self.drawRightString(w - 12 * mm, h - 9.5 * mm, date_txt)
            self.setStrokeColor(LINE)
            self.setLineWidth(0.5)
            self.line(12 * mm, 11 * mm, w - 12 * mm, 11 * mm)
            self.setFillColor(MUTED)
            self.setFont("Helvetica", 7.5)
            self.drawString(12 * mm, 7 * mm, f"Generated {generated}")
            self.drawRightString(w - 12 * mm, 7 * mm, f"Page {self._pageNumber} of {total}")

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                          topMargin=24 * mm, bottomMargin=15 * mm, title=f"{title} {iso}", author=ORG_NAME)
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(
        doc.leftMargin, doc.bottomMargin, doc.width, doc.height,
        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)])])
    W = doc.width

    def ps(name, **kw):
        base = dict(fontName="Helvetica", fontSize=8.3, leading=10.2, textColor=INK)
        base.update(kw)
        return ParagraphStyle(name, **base)

    cell, cellb = ps("cell"), ps("cellb", fontName="Helvetica-Bold")
    head = ps("head", fontName="Helvetica-Bold", fontSize=7.6, textColor=colors.white)
    muted = ps("muted", textColor=MUTED, fontSize=8)
    P = lambda t, st=cell: Paragraph(escape(_s(t)) or "&nbsp;", st)  # noqa: E731

    story = []
    col_w = [9, 46, 36, 24, 27, 70, 34, 27]
    col_w = [x * mm * W / (sum(col_w) * mm) for x in col_w]

    def duty_table(df, band_text, bg, fg, start_no, off=False):
        band = Paragraph(f"<b>{escape(band_text)}</b>", ps("band", fontSize=9.5, leading=12, textColor=fg))
        labels = ["#", "NAME", "DESIGNATION", "BELT NO.", "CONTACT",
                  "STATUS" if off else "BEAT / AREA", "" if off else "TIME", "REMARKS"]
        data = [[band] + [""] * 7, [Paragraph(x, head) for x in labels]]
        style = [
            ("SPAN", (0, 0), (-1, 0)), ("BACKGROUND", (0, 0), (-1, 0), bg),
            ("BACKGROUND", (0, 1), (-1, 1), NAVY), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.6, LINE), ("LINEBELOW", (0, 2), (-1, -1), 0.25, LINE),
            ("TOPPADDING", (0, 0), (-1, -1), 3.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.2),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("NOSPLIT", (0, 0), (-1, 2)),
        ]
        group, last = -1, None
        for i, r in enumerate(df.itertuples(index=False), start=2):
            key = (r.beat, r.slot) if not off else i
            if key != last:
                group, last = group + 1, key
            if group % 2:
                style.append(("BACKGROUND", (0, i), (-1, i), ZEBRA))
            data.append([
                P(start_no + i - 2, ps("n", textColor=MUTED)), P(r.name, cellb), P(r.designation), P(r.belt_no),
                P(r.contact), P(r.status if off else r.beat), P("" if off else r.slot), P(r.remarks, ps("rm", fontSize=7.6)),
            ])
        t = Table(data, colWidths=col_w, repeatRows=2)
        t.setStyle(style)
        return t

    def kpi_strip(vals):
        tints = ["#E8F1FF", "#FFF3D1", "#DDF3E4", "#EEF1F6"]
        cells = [[Paragraph(
            f"<font size=17 color='#12263F'><b>{v}</b></font><br/><font size=7 color='#64748B'>{lab}</font>",
            ps("k", leading=19, alignment=1)) for lab, v in vals]]
        t = Table(cells, colWidths=[W / len(vals)] * len(vals))
        t.setStyle([("BACKGROUND", (i, 0), (i, 0), colors.HexColor(tints[i % 4])) for i in range(len(vals))] + [
            ("LINEAFTER", (0, 0), (-2, -1), 3, colors.white), ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")])
        return t

    def signature_block(station):
        sig = get_signature(station)
        if sig["image"]:
            from PIL import Image as PILImage
            iw, ih = PILImage.open(io.BytesIO(sig["image"])).size
            scale = min(52 * mm / iw, 15 * mm / ih)
            top = Image(io.BytesIO(sig["image"]), width=iw * scale, height=ih * scale)
        else:
            top = Spacer(1, 15 * mm)
        rows_ = [[top], [Paragraph(f"<b>{escape(sig['label'])}</b>", ps("sg", alignment=1, fontSize=9.5, leading=12))]]
        if sig["officer"]:
            rows_.append([Paragraph(escape(sig["officer"]), ps("so", alignment=1, textColor=MUTED))])
        t = Table(rows_, colWidths=[75 * mm], hAlign="RIGHT")
        t.setStyle([("LINEABOVE", (0, 1), (-1, 1), 0.8, INK), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, 0), "BOTTOM"), ("TOPPADDING", (0, 1), (-1, 1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 2)])
        return t

    # ---- cover page (consolidated only) -------------------------------------
    if cover:
        story.append(Marker("ALL STATIONS"))
        story.append(Paragraph("<b>Station-wise summary</b>", ps("h", fontSize=13, leading=16)))
        story.append(Spacer(1, 3 * mm))
        hdr = ["STATION", "TOTAL STAFF", "ON DUTY", "LEAVE / OFF", "UNASSIGNED", "BEATS COVERED", "ROSTER STATUS"]
        widths = [W * x for x in (0.26, 0.11, 0.11, 0.11, 0.11, 0.13, 0.17)]
        data = [[Paragraph(x, head) for x in hdr]]
        status_txt = {"final": ("FINAL", "#17653A"), "draft": ("DRAFT - not marked final", "#B45309"),
                      "none": ("NOT CREATED", "#B91C1C")}
        for r in summ.itertuples(index=False):
            txt, col = status_txt[r.state]
            data.append([P(r.station, cellb), P(r.staff), P(r.on_duty, cellb), P(r.leave_off), P(r.unassigned),
                         P(r.beats), Paragraph(f"<font color='{col}'><b>{txt}</b></font>", cell)])
        data.append([P("TOTAL", cellb)] + [P(int(summ[c].sum()), cellb) for c in
                                           ("staff", "on_duty", "leave_off", "unassigned", "beats")] + [P("")])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3EAF3")),
                    ("BOX", (0, 0), (-1, -1), 0.6, LINE), ("LINEBELOW", (0, 1), (-1, -1), 0.25, LINE),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 6)] +
                   [("BACKGROUND", (0, i), (-1, i), ZEBRA) for i in range(2, len(data) - 1, 2)])
        story += [t, Spacer(1, 5 * mm),
                  Paragraph("Station rosters follow, one station per page group, each with the station's own "
                            "SDEO signature. UNASSIGNED = active staff with no duty or leave entry for the day.", muted)]

    # ---- one section per station ---------------------------------------------
    printed = 0
    for stn in stations:
        r = rows[rows["station"] == stn]
        if r.empty and cover:
            continue
        if cover or printed:
            story.append(PageBreak())
        printed += 1
        story.append(Marker(stn))
        duty, off = r[r["status"] == "Duty"], r[r["status"] != "Duty"]
        srow = summ[summ["station"] == stn].iloc[0]
        story.append(kpi_strip([("ON DUTY", len(duty)), ("LEAVE / OFF / TRAINING", len(off)),
                                ("BEATS COVERED", int(srow["beats"])), ("TOTAL STAFF", int(srow["staff"]))]))
        story.append(Spacer(1, 4 * mm))
        if r.empty:
            story.append(Paragraph("No duties have been assigned for this date.", ps("e", fontSize=11, leading=14)))
        no = 1
        for shift, g in duty.groupby(duty["shift"].replace("", "Shift not set"), sort=False):
            bg, fg = band_colors(shift)
            slots = list(dict.fromkeys(x for x in g["slot"] if x))
            note = f"   |   {slots[0]}" if len(slots) == 1 else ""
            story.append(duty_table(g, f"{shift.upper()}{note}   |   {len(g)} staff", bg, fg, no))
            story.append(Spacer(1, 4 * mm))
            no += len(g)
        if not off.empty:
            story.append(duty_table(off, f"NOT ON DUTY   |   {len(off)} staff", colors.HexColor("#F1F5F9"),
                                    colors.HexColor("#475569"), 1, off=True))
            story.append(Spacer(1, 4 * mm))
        story += [CondPageBreak(42 * mm), Spacer(1, 6 * mm), signature_block(stn)]

    if not stations or (not cover and printed == 0):
        story = [Marker("-"), Paragraph("No station selected.", cell)]
    doc.build(story, canvasmaker=Numbered)
    return buf.getvalue()


# =============================================================================
# EXPORT: EXCEL
# =============================================================================
def build_excel(iso: str, stations) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = roster_rows(iso, stations)
    summ = day_summary(iso, stations)
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    head = ["Station", "Date", "Sr", "Name", "Designation", "Belt No.", "Contact", "Status",
            "Beat / Area", "Shift", "Time Slot", "Remarks"]
    ws.append(head)
    for i, r in enumerate(rows.itertuples(index=False), start=1):
        ws.append([nice(r.station), iso, i, r.name, r.designation, r.belt_no, r.contact, r.status,
                   r.beat, r.shift, r.slot, r.remarks])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Station", "Total Staff", "On Duty", "Leave / Off", "Unassigned", "Beats Covered", "Roster Status"])
    for r in summ.itertuples(index=False):
        ws2.append([nice(r.station), r.staff, r.on_duty, r.leave_off, r.unassigned, r.beats,
                    {"final": "Final", "draft": "Draft", "none": "Not created"}[r.state]])
    for sheet in (ws, ws2):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12263F")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for i, col in enumerate(sheet.columns, start=1):
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =============================================================================
# BACKUP / RESTORE / AUDIT
# =============================================================================
def backup_bytes() -> bytes:
    fd, tmp = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        src, dst = sqlite3.connect(DB_FILE), sqlite3.connect(tmp)
        src.backup(dst)
        src.close()
        dst.close()
        with open(tmp, "rb") as f:
            return f.read()
    finally:
        os.remove(tmp)


def restore_backup(data: bytes):
    fd, tmp = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        with open(tmp, "wb") as f:
            f.write(data)
        try:
            chk = sqlite3.connect(tmp)
            tables = {r[0] for r in chk.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            chk.close()
        except sqlite3.DatabaseError:
            return False, "That file is not a valid backup."
        if not {"users", "staff", "roster", "lists", "slots"} <= tables:
            return False, "That file is not a backup of this app."
        shutil.copyfile(tmp, DB_FILE)
    finally:
        os.remove(tmp)
    init_db()
    return True, "Backup restored."


def get_audit(limit: int = 300) -> pd.DataFrame:
    with db() as c:
        return pd.read_sql_query("SELECT ts AS time, user, action, detail FROM audit ORDER BY id DESC LIMIT ?",
                                 c, params=(limit,))


# =============================================================================
# FIRST-RUN: schema, accounts, sample data
# =============================================================================
def _seed_from_excel(c):
    df = pd.read_excel(SEED_FILE, sheet_name=0, dtype=str)
    need = {"Station", "Name", "Belt No.", "Designation", "Contact No.", "Squad / Team", "Beat Area",
            "Shift", "Start Time", "End Time"}
    if not need <= set(df.columns):
        return
    placed = set()
    for r in df.to_dict("records"):
        stn = " ".join(_s(r["Station"]).split()).upper()
        name = norm_name(r["Name"])
        if not stn or not name:
            continue
        if not c.execute("SELECT 1 FROM users WHERE username=?", (stn,)).fetchone():
            _set_pw(c, stn, "station", secrets.token_urlsafe(9), True)
        beat, shift = " ".join(_s(r["Beat Area"]).split()), " ".join(_s(r["Shift"]).split())
        s, e = to_hhmm(r["Start Time"]), to_hhmm(r["End Time"])
        slot = slot_label(s, e) if s and e else ""
        _add_list(c, "beat", stn, [beat])
        _add_list(c, "shift", stn, [shift])
        if slot:
            c.execute("INSERT OR IGNORE INTO slots(station, t_start, t_end, label) VALUES (?,?,?,?)", (stn, s, e, slot))
        belt = norm_belt(r["Belt No."])
        key = (stn, belt or name.lower())
        if key in placed:
            continue
        placed.add(key)
        sid = c.execute("SELECT id FROM staff WHERE belt_no=? AND belt_no<>''", (belt,)).fetchone()
        sid = sid[0] if sid else _insert_staff(c, stn, name, norm_designation(r["Designation"]), belt,
                                                norm_phone(r["Contact No."]))
        if sid and SEED_ROSTER:
            squad = " ".join(_s(r["Squad / Team"]).split())
            _upsert(c, stn, today_iso(), sid, "Duty", beat, shift, slot, squad, "seed")


def _run_schema(conn) -> None:
    """Run SCHEMA one statement at a time so a failure names the exact statement."""
    buf = ""
    for line in SCHEMA.splitlines(keepends=True):
        buf += line
        if sqlite3.complete_statement(buf):
            try:
                conn.execute(buf)
            except sqlite3.Error as exc:
                raise sqlite3.OperationalError(
                    f"{exc} | SQLite {sqlite3.sqlite_version} | statement: {' '.join(buf.split())[:160]}") from exc
            buf = ""


def diagnostics() -> str:
    """Plain-text facts for the on-screen error box (the host hides the real exception)."""
    try:
        size = os.path.getsize(DB_FILE) if os.path.exists(DB_FILE) else None
    except OSError:
        size = "?"
    return (f"SQLite {sqlite3.sqlite_version} | db={DB_FILE} | exists_size={size} | "
            f"dir_writable={os.access(os.path.dirname(DB_FILE) or '.', os.W_OK)} | "
            f"core.py={os.path.abspath(__file__)}")


def _expected_columns() -> dict:
    """Tables and columns that SCHEMA defines, read from a throw-away in-memory database."""
    m = sqlite3.connect(":memory:")
    _run_schema(m)
    out = {t: {r[1] for r in m.execute(f"PRAGMA table_info({t})")}
           for (t,) in m.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
    m.close()
    return out


def _move_aside(reason: str) -> None:
    dest = f"{DB_FILE}.legacy-{datetime.now():%Y%m%d%H%M%S}"
    print(f"Database set aside ({reason}) -> {dest}")
    shutil.move(DB_FILE, dest)


def _set_aside_legacy_db():
    """An old-format or incompatible database is renamed, never overwritten."""
    if not os.path.exists(DB_FILE):
        return
    try:
        chk = sqlite3.connect(DB_FILE)
        tables = {r[0] for r in chk.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        found = {t: {r[1] for r in chk.execute(f"PRAGMA table_info({t})")} for t in tables}
        chk.close()
    except sqlite3.DatabaseError:
        _move_aside("unreadable file")
        return
    cols = found.get("users", set())
    if "employees" in tables or ("users" in tables and "role" not in cols):
        _move_aside("old app version")
        return
    for t, want in _expected_columns().items():      # same table name, different columns
        if t in found and not want <= found[t]:
            _move_aside(f"table '{t}' is missing {sorted(want - found[t])}")
            return


def init_db() -> bool:
    """Start-up with self-repair: a database the app cannot use is set aside and rebuilt,
    and an unwritable folder (some cloud hosts) falls back to the system temp folder."""
    global DB_FILE
    for attempt in range(3):
        try:
            _set_aside_legacy_db()
            return _init_db_once()
        except sqlite3.OperationalError as exc:
            msg = str(exc).lower()
            print(f"init_db attempt {attempt + 1} failed: {exc} (db={DB_FILE})")
            if attempt == 0 and any(k in msg for k in ("readonly", "unable to open", "disk i/o")):
                DB_FILE = os.path.join(tempfile.gettempdir(), "duty_roster_v3.db")
            elif os.path.exists(DB_FILE):
                _move_aside(str(exc))
            else:
                raise
    raise RuntimeError("Could not initialise the database")


def _init_db_once() -> bool:
    with db() as c:
        _run_schema(c)
        if not c.execute("SELECT 1 FROM users WHERE role='admin'").fetchone():
            _set_pw(c, ADMIN_USER, "admin", DEFAULT_ADMIN_PW, True)
        if not c.execute("SELECT 1 FROM users WHERE role='station'").fetchone():
            for stn, pw in DEFAULT_STATIONS.items():
                _set_pw(c, stn, "station", pw, True)
            if os.path.exists(SEED_FILE):
                try:
                    _seed_from_excel(c)
                except Exception as exc:  # noqa: BLE001 - a bad sample file must never block start-up
                    print("Seed skipped:", exc)
        for (stn,) in c.execute("SELECT username FROM users WHERE role='station'").fetchall():
            _ensure_defaults(c, stn)
    return True